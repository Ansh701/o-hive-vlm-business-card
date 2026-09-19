from __future__ import annotations

from io import BytesIO
from uuid import UUID

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from backend.app.config import Settings
from backend.app.export import COLUMNS, build_workbook, export_filename
from backend.app.inference import BusinessCardLead
from backend.app.main import create_app
from backend.app.models import Lead, LeadStatus


def lead(**changes: object) -> Lead:
    values: dict[str, object] = {
        "batch_id": UUID("3d7dbeb6-c61b-45a7-b058-0f8058740d6b"),
        "source_filename": "card.png",
        "content_sha256": "a" * 64,
        "first_name": "Ada",
        "last_name": "Lovelace",
        "job_title": "Engineer",
        "company": "Difference Works",
        "location": "London",
        "phone_number": "+442079460000",
        "email": "ada@example.com",
        "status": LeadStatus.SUCCESS,
    }
    values.update(changes)
    return Lead(**values)


def workbook_rows(payload: bytes) -> tuple[object, list[tuple[object, ...]]]:
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = workbook.active
    return sheet, list(sheet.iter_rows(values_only=True))


def test_workbook_is_genuine_xlsx_with_exact_required_column_order() -> None:
    payload = build_workbook([lead()]).getvalue()
    sheet, rows = workbook_rows(payload)

    assert payload.startswith(b"PK")
    assert rows[0] == tuple(label for label, _field in COLUMNS)
    assert rows[0] == (
        "First Name",
        "Last Name",
        "Position / Job Title",
        "Company",
        "Location",
        "Phone Number",
        "Email Address",
    )
    assert rows[1] == (
        "Ada",
        "Lovelace",
        "Engineer",
        "Difference Works",
        "London",
        "'+442079460000",
        "ada@example.com",
    )
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G2"
    assert sheet["A1"].font.bold is True


@pytest.mark.parametrize("dangerous", ["=1+1", "+SUM(A1:A2)", "-2+3", "@cmd", "\t=cmd"])
def test_formula_leading_values_are_exported_as_inert_text(dangerous: str) -> None:
    _sheet, rows = workbook_rows(build_workbook([lead(company=dangerous)]).getvalue())

    assert rows[1][3] == f"'{dangerous}"


def test_missing_fields_export_as_empty_cells() -> None:
    _sheet, rows = workbook_rows(
        build_workbook([lead(last_name=None, location=None, email=None)]).getvalue()
    )

    assert rows[1][1] is None
    assert rows[1][4] is None
    assert rows[1][6] is None


def test_column_width_is_bounded_for_hostile_long_text() -> None:
    sheet, _rows = workbook_rows(build_workbook([lead(company="x" * 300)]).getvalue())

    assert 12 <= sheet.column_dimensions["D"].width <= 50


def test_export_filename_is_sanitized_and_dated() -> None:
    assert export_filename(date_text="2026-09-19") == "business-card-leads-2026-09-19.xlsx"


class ExportInference:
    async def extract(self, _image: bytes, _media_type: str) -> BusinessCardLead:
        return BusinessCardLead(first_name="Initial", company="Safe")


def png_bytes() -> bytes:
    from PIL import Image

    buffer = BytesIO()
    Image.new("RGB", (640, 360), "white").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_export_endpoint_uses_selected_corrected_leads_only(
    session_factory: async_sessionmaker[AsyncSession],
) -> None:
    app = create_app(
        settings=Settings(environment="test"),
        session_factory=session_factory,
        inference_client=ExportInference(),
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://testserver"
    ) as client:
        batch_id = (await client.post("/api/batches", json={"total_cards": 2})).json()["id"]
        first = await client.post(
            f"/api/batches/{batch_id}/cards",
            files={"file": ("one.png", png_bytes(), "image/png")},
        )
        from PIL import Image

        other_buffer = BytesIO()
        Image.new("RGB", (640, 360), "black").save(other_buffer, format="PNG")
        second = await client.post(
            f"/api/batches/{batch_id}/cards",
            files={"file": ("two.png", other_buffer.getvalue(), "image/png")},
        )
        first_id = first.json()["id"]
        second_id = second.json()["id"]
        await client.patch(
            f"/api/leads/{first_id}", json={"first_name": "Corrected", "company": "=BAD()"}
        )
        response = await client.get(
            f"/api/batches/{batch_id}/export.xlsx", params=[("lead_ids", first_id)]
        )

    _sheet, rows = workbook_rows(response.content)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "business-card-leads-" in response.headers["content-disposition"]
    assert len(rows) == 2
    assert rows[1][0] == "Corrected"
    assert rows[1][3] == "'=BAD()"
    assert second_id not in response.text


@pytest.mark.asyncio
async def test_export_rejects_empty_selection_and_foreign_lead(
    session_factory: async_sessionmaker[AsyncSession],
) -> None:
    app = create_app(
        settings=Settings(environment="test"),
        session_factory=session_factory,
        inference_client=ExportInference(),
    )
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://testserver"
    ) as client:
        first_batch = (await client.post("/api/batches", json={"total_cards": 1})).json()["id"]
        first_lead = await client.post(
            f"/api/batches/{first_batch}/cards",
            files={"file": ("one.png", png_bytes(), "image/png")},
        )
        second_batch = (await client.post("/api/batches", json={"total_cards": 1})).json()["id"]
        empty = await client.get(f"/api/batches/{second_batch}/export.xlsx")
        foreign = await client.get(
            f"/api/batches/{second_batch}/export.xlsx",
            params=[("lead_ids", first_lead.json()["id"])],
        )

    assert empty.status_code == 422
    assert empty.json()["code"] == "no_exportable_leads"
    assert foreign.status_code == 422
    assert foreign.json()["code"] == "no_exportable_leads"
