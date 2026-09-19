from __future__ import annotations

from collections.abc import Iterable
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.app.models import Lead

COLUMNS = (
    ("First Name", "first_name"),
    ("Last Name", "last_name"),
    ("Position / Job Title", "job_title"),
    ("Company", "company"),
    ("Location", "location"),
    ("Phone Number", "phone_number"),
    ("Email Address", "email"),
)


def _spreadsheet_safe(value: str | None) -> str | None:
    """Keep model/user text from being interpreted as a spreadsheet formula."""
    if value is None:
        return None
    first_visible = value.lstrip(" \t\r\n")[:1]
    if first_visible in {"=", "+", "-", "@"}:
        return f"'{value}"
    return value


def build_workbook(leads: Iterable[Lead]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"

    headers = [label for label, _field in COLUMNS]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="173F3A")

    row_count = 1
    for lead in leads:
        sheet.append(
            [_spreadsheet_safe(getattr(lead, field)) for _label, field in COLUMNS]
        )
        row_count += 1

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:G{row_count}"
    sheet.sheet_view.showGridLines = False
    for index, (header, _field) in enumerate(COLUMNS, start=1):
        contents = [header]
        contents.extend(
            str(sheet.cell(row=row, column=index).value or "")
            for row in range(2, row_count + 1)
        )
        width = min(50, max(12, max(len(value) for value in contents) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_filename(*, date_text: str | None = None) -> str:
    safe_date = date_text or datetime.now(UTC).date().isoformat()
    return f"business-card-leads-{safe_date}.xlsx"
